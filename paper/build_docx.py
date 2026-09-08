#!/usr/bin/env python3
"""Build an Elsevier-style single-column, double-spaced Word manuscript."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "KBS_short_communication.docx"
FIG1 = ROOT / "Figure_1.png"
FIG2 = ROOT / "Figure_2.png"


def set_run_font(run, size=12, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_paragraph_format(p, *, after=0, before=0, first_line=True, align="left", space=2.0):
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = space
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.first_line_indent = Pt(18) if first_line else Pt(0)


def add_text(doc, text, *, size=12, bold=False, italic=False, after=0, before=0,
             first_line=True, align="justify"):
    p = doc.add_paragraph()
    set_paragraph_format(p, after=after, before=before, first_line=first_line, align=align)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_h1(doc, text):
    p = doc.add_heading(text, level=1)
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 2.0
    pf.space_after = Pt(0)
    pf.space_before = Pt(12)
    for run in p.runs:
        set_run_font(run, size=12, bold=True)
    return p


def add_h2(doc, text):
    p = doc.add_heading(text, level=2)
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 2.0
    pf.space_after = Pt(0)
    pf.space_before = Pt(8)
    for run in p.runs:
        set_run_font(run, size=12, bold=True)
    return p


def shade_header(cell):
    tc = cell._tePr if hasattr(cell, "_tePr") else cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "D9E2F3")
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "666666")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def fill_table(doc, headers, rows, col_widths):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        set_run_font(run, size=8, bold=True)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        shade_header(cell)
        set_cell_border(cell)
        cell.width = Inches(col_widths[i])
    for r_i, row in enumerate(rows, start=1):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i].cells[c_i]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_i else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(val))
            set_run_font(run, size=8)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            set_cell_border(cell)
            cell.width = Inches(col_widths[c_i])
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, after=12, before=6, first_line=False, align="left", space=1.15)
    run = p.add_run(text)
    set_run_font(run, size=11, italic=True)


def add_picture(doc, path, width=6.3):
    p = doc.add_paragraph()
    set_paragraph_format(p, after=0, before=6, first_line=False, align="center", space=1.0)
    p.add_run().add_picture(str(path), width=Inches(width))


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 2.0

    add_text(
        doc,
        "Knowledge-Guided Decision Intelligence for Cost-Aware Customer Churn Intervention Under Predictive Uncertainty",
        size=14, bold=True, first_line=False, align="center", after=6,
    )
    add_text(doc, "Short Communication", size=12, first_line=False, align="center", after=12)
    add_text(doc, "Dr. Sai Krishna Thota, PhD", size=12, first_line=False, align="center", after=0)
    add_text(
        doc,
        "Independent Researcher, USA",
        size=12, first_line=False, align="center", after=18,
    )

    add_h1(doc, "Abstract")
    add_text(
        doc,
        "Customer churn scores are often treated as if they were already an intervention policy. Existing studies largely optimise prediction or expected profit, while calibrated risk, structured domain knowledge, strategic customer value, predictive uncertainty, human review, and a shared budget are rarely treated as one constrained decision problem. This short communication places a leakage-safe calibrated ensemble under a knowledge-guided allocation layer. A service-contract graph supplies out-of-fold motif risk. IBM customer lifetime value (CLTV) is a dimensionless strategic-priority index; simulated annual margin from monthly charges is the economic value in the utility. Under a common expected-cost budget, most spend goes to positive expected-utility offers and a minority is reserved for human review of high-CLTV cases. On IBM Telco data (N = 7043), ensemble AUROC is 0.860 ± 0.001 and expected calibration error is 0.012 ± 0.003. The prespecified primary decision endpoint is high-value churn reach. Versus cost-aware targeting it rises by 12.6 percentage points (1000/1000 paired bootstrap resamples positive; 95% interval [9.3, 16.5] pp). Overall churn reach stays comparable. The net-benefit proxy rises by $2320 (968/1000 resamples positive; 95% interval [-$92, $5397]) and is not claimed as confirmed economic superiority. Component ablation shows that this coverage gain is produced by the reserved high-CLTV review slice, not by incremental graph-residual routing at the nominal 85/15 operating point. The contribution is a constrained decision-support procedure, not a new classifier.",
        first_line=False,
    )
    p = doc.add_paragraph()
    set_paragraph_format(p, after=0, before=6, first_line=False, align="left")
    r0 = p.add_run("Keywords: ")
    set_run_font(r0, size=12, bold=True)
    r1 = p.add_run("churn prediction; decision support; knowledge graph; predictive uncertainty; cost-sensitive learning; human-in-the-loop")
    set_run_font(r1, size=12, bold=False)

    add_h1(doc, "1. Introduction")
    add_text(
        doc,
        "Recent surveys of machine-learning churn research show that predictive pipelines are mature, yet profit-oriented evaluation, operational allocation, and explainable decision support remain comparatively thin [1,2]. Ensemble and boosting classifiers dominate published benchmarks [2,9,10]. Profit-driven ensembles and prescriptive methods that trade expected profit against discrimination, including heterogeneous lifetime value and incentive cost, improve monetary scores relative to accuracy-only training [4,5]. Probability calibration is required if those scores are to be used as decision inputs [11]. Human inspection belongs inside a decision-support system rather than as an afterthought [12]. Knowledge bases and graphs have been used to enrich churn features [3] and, in other domains, to support structured decision-making [13,14].",
    )
    add_text(
        doc,
        "The scientific gap is more specific than better churn AUROC. Existing work largely optimises prediction or expected profit in isolation. Comparatively little work treats calibrated risk, structured domain knowledge, strategic customer value, predictive uncertainty, scarce human review, and a shared intervention budget as one constrained decision problem. Two operational facts make that joint problem necessary. First, economic targeting based on monthly charges is not equivalent to strategic targeting based on lifetime-value knowledge: in the IBM Telco workbook, CLTV and the simulated annual-margin proxy have Pearson correlation 0.10. Second, converting predictive entropy into a blanket review rule can consume the budget with cheap reviews and leave too little capacity for high-utility offers.",
    )
    add_text(
        doc,
        "This short communication therefore separates prediction from allocation. A calibrated ensemble estimates churn probability. A compact service-contract graph and CLTV encode domain structure. The proposed knowledge-guided decision intelligence (KGDI) policy is budget-matched: majority economic allocation plus a reserved review slice for high-CLTV customers. High-value churn reach under that shared budget is the primary decision endpoint. Overall churn reach and an outcome-anchored net-benefit proxy are secondary. The 12.6 percentage-point coverage gain is interpreted as achievement of that intended strategic objective, not as unexpected discovery of a superior classifier. Only two tables and two figures are reported.",
    )

    add_h1(doc, "2. Proposed Framework")
    add_h2(doc, "2.1 Predictive model")
    add_text(
        doc,
        "Leakage-safe repeated stratified out-of-fold prediction is used (5 folds × 5 repeats). Predictors are an allowlist of service, contract, billing, and tenure fields. Direct leakage fields (churn label, churn score, churn reason) are excluded. Gender and senior-citizen status are excluded from the model and from the policy. Logistic regression, random forest, XGBoost, and LightGBM are calibrated with a sigmoid map fitted on a dedicated calibration split inside each training fold [11]. The ensemble probability pᵢ is the equal-weight mean of the four calibrated probabilities. Primary uncertainty is the normalised binary entropy of pᵢ. Between-model disagreement is retained only as a diagnostic. No synthetic oversampling is applied; imbalance is handled through cost-aware utilities [6-8]. Ranking and calibration, not F1, feed Eqs. (1)-(2).",
    )

    add_h2(doc, "2.2 Knowledge representation")
    add_text(
        doc,
        "A heterogeneous service-contract graph is built from observed IBM fields. Customer nodes link to contract type, internet service, payment method, and add-on or protection services. Motifs were specified from service-contract semantics that are standard in telecommunications operations-month-to-month fibre, electronic-check payment, unprotected internet, and short tenure-before fold-local rate estimation. They were not selected by scanning outcome labels on the full file. One-hop knowledge risk kᵢ for a held-out customer is the training-fold empirical churn rate of incident motifs, averaged out-of-fold. Test-fold labels never enter motif rates. Knowledge risk is a structural prior (AUROC 0.81) rather than a replacement for the ensemble (AUROC 0.86; correlation of kᵢ with pᵢ = 0.76). The residual kᵢ − pᵢ is available as an optional routing flag. This construction is a domain graph for decision gating, not an industrial ontology, a large language-model knowledge base, or a graph-reasoning engine [3,13,14]. Related work on assembling and reusing structured relations in other domains is cited only for that limited precedent [15-20].",
    )
    add_text(
        doc,
        "IBM CLTV is a customer-node attribute converted to a percentile rank and is never treated as currency. Economic value in the utility is the simulated margin proxy Vᵢ = monthly charges × 12 × 0.60.",
    )

    add_h2(doc, "2.3 Decision model")
    add_text(
        doc,
        "Let cₒ be offer cost, s the assumed retention-success probability, cₕ the review cost, and (Se, Sp) reviewer sensitivity and specificity. Expected offer and review utilities are",
        first_line=True,
    )
    add_text(doc, "Uᵢ^offer = pᵢ s Vᵢ − cₒ,                                            (1)", first_line=False, align="center")
    add_text(
        doc,
        "Uᵢ^review = pᵢ Se s Vᵢ − (cₕ + cₒ [pᵢ Se + (1 − pᵢ)(1 − Sp)]).     (2)",
        first_line=False, align="center",
    )
    add_text(
        doc,
        "The planning budget is B = N cₒ × 0.20 and is enforced on ex-ante expected selected cost. Three baselines receive the same B: (i) probability-threshold ranking among pᵢ ≥ 0.5; (ii) cost-aware greedy ranking by Uᵢ^offer among pᵢ ≥ 0.20 and Uᵢ^offer > 0; (iii) uncertainty review, which routes high-entropy eligible customers to human review and otherwise offers [4,5,12].",
    )
    add_text(
        doc,
        "The proposed KGDI policy uses a nominal operating point of 85% of B on cost-aware offers and 15% reserved for human review. That 15% represents a deliberately limited minority review capacity, not a unique optimum. A robustness path over economic shares {0.70, 0.75, 0.80, 0.85, 0.90, 1.00} is shown in Fig. 2. Review eligibility requires pᵢ ≥ 0.20, CLTV at or above the 75th percentile, positive review utility, and either moderate-or-higher entropy or kᵢ − pᵢ > 0. Leftover budget returns to cost-aware offers. Observed churn is not used to choose actions.",
    )
    add_text(
        doc,
        "Default values are cₒ = $100, s = 0.35, cₕ = $25, Se = Sp = 0.85. Inference versus cost-aware targeting uses 1000 paired customer bootstrap resamples with policy rerouting and budget reallocation inside every resample. The bootstrap conditions on the previously generated out-of-fold risk estimates and therefore quantifies decision-allocation uncertainty rather than full model-training uncertainty. Reported quantities are the observed delta, a percentile 95% interval, and the fraction of resamples with a positive delta. That fraction is not a classical p-value.",
    )

    add_h1(doc, "3. Experimental Design")
    add_text(
        doc,
        "The IBM Telco Customer Churn workbook contains 7043 customers and a 26.5% churn rate. High-value churners, used only for evaluation, are observed churners with CLTV percentile at least 0.75 (n = 365). Because the proposed policy also uses the 0.75 CLTV quantile as a review gate, high-value reach is an intended strategic coverage objective, not an independent surprise metric. All policies share planning budget $140,860. Public telecommunications tables of this type remain a standard test-bed [1,2,7]. Primary decision endpoint: high-value churn reach under the common expected-cost budget. Secondary decision endpoints: overall churn reach and the outcome-anchored net-benefit proxy. Component ablations and scenario perturbations are reported in the Results as supplementary analyses, not as additional primary tables. Intervention rate may exceed 20% when cheaper reviews replace offers while expected selected cost remains at B.",
    )

    add_h1(doc, "4. Results and Discussion")
    add_text(
        doc,
        "Table 1 shows that the ensemble is well calibrated (Brier 0.130 ± 0.001; expected calibration error 0.012 ± 0.003). Table 1 licences the probabilities in Eqs. (1)-(2). No predictive state of the art is claimed [1,2,7,9]. Logistic regression has a higher F1 than the ensemble; the policy uses ranking and calibration rather than F1.",
    )

    add_caption(
        doc,
        "Table 1. Repeated out-of-fold predictive performance (mean ± sd over five repeats).",
    )
    fill_table(
        doc,
        ["Model", "AUROC", "PR-AUC", "F1", "Brier ↓", "ECE ↓"],
        [
            ["Logistic regression", "0.856 ± 0.000", "0.669 ± 0.002", "0.616 ± 0.003", "0.132 ± 0.000", "0.013 ± 0.003"],
            ["Random forest", "0.849 ± 0.002", "0.656 ± 0.008", "0.563 ± 0.011", "0.135 ± 0.001", "0.015 ± 0.002"],
            ["XGBoost", "0.859 ± 0.001", "0.679 ± 0.005", "0.600 ± 0.006", "0.131 ± 0.000", "0.011 ± 0.003"],
            ["LightGBM", "0.851 ± 0.001", "0.665 ± 0.005", "0.586 ± 0.006", "0.134 ± 0.000", "0.014 ± 0.003"],
            ["Ensemble", "0.860 ± 0.001", "0.680 ± 0.005", "0.601 ± 0.005", "0.130 ± 0.001", "0.012 ± 0.003"],
        ],
        [1.6, 1.15, 1.15, 1.15, 1.15, 1.15],
    )
    add_text(
        doc,
        "Notes: Brier and ECE are lower-is-better. ECE: expected calibration error.",
        size=10, italic=True, first_line=False, after=12,
    )

    add_text(
        doc,
        "Table 2 is the decision result. Cost-aware targeting is the strongest economic baseline among the offer-only policies ($65,209). Uncertainty review raises high-value reach to 54.2% with 1650 reviews, but the proxy falls to $59,253. Proposed KGDI issues 1297 offers and 166 reviews (intervention rate 20.8% because reviews are cheaper than offers; expected cost remains at B). High-value churn reach is 56.4% versus 43.8% for cost-aware targeting (+12.6 percentage points; 1000/1000 resamples positive; 95% interval [9.3, 16.5] pp). Overall churn reach is essentially unchanged (50.6% versus 49.9%). The net-benefit proxy is $67,529 (+$2320). In 968/1000 resamples the economic delta was positive, but the 95% interval [-$92, $5397] includes a small negative value, so economic dominance is not confirmed.",
    )

    add_caption(
        doc,
        "Table 2. Budget-matched decision performance. Deltas are versus cost-aware targeting (1000 paired policy-rerun bootstrap resamples).",
    )
    fill_table(
        doc,
        [
            "Strategy", "Interv. %", "Churn reach %", "HV reach %",
            "Offers / reviews", "Net-benefit (USD)", "Δ USD", "95% CI",
            "P(Δ>0) $", "Δ HV (pp)", "P(Δ>0) HV",
        ],
        [
            ["Probability threshold", "20.0", "51.7", "43.0", "1408 / 0", "53,786", "—", "—", "—", "—", "—"],
            ["Cost-aware", "20.0", "49.9", "43.8", "1408 / 0", "65,209", "—", "—", "—", "—", "—"],
            ["Uncertainty review", "24.5", "56.4", "54.2", "79 / 1650", "59,253", "—", "—", "—", "—", "—"],
            ["Proposed KGDI", "20.8", "50.6", "56.4", "1297 / 166", "67,529", "2,320", "[-92, 5397]", "968/1000", "12.6", "1000/1000"],
        ],
        [1.15, 0.55, 0.62, 0.58, 0.85, 0.85, 0.48, 0.78, 0.55, 0.58, 0.58],
    )
    add_text(
        doc,
        "Notes: Interv.: intervention rate. HV: high-value churn reach. CI: percentile interval. Resample fractions are not classical p-values. High-value Δ 95% interval [9.3, 16.5] pp.",
        size=10, italic=True, first_line=False, after=12,
    )

    add_text(
        doc,
        "Component ablation under the same 85/15 split shows why the coverage gain occurs. A CLTV-only review reserve (no entropy filter, no graph residual) produces the same 1297 offers, 166 reviews, 50.6% overall reach, 56.4% high-value reach, and $67,529 proxy as full KGDI (0 of 7043 actions differ). Adding entropy does not change that set. CLTV plus graph residual without entropy yields fewer reviews (101) and lower high-value reach (51.8%). Removing the CLTV gate increases reviews to 305 and drops high-value reach to 47.1%. Setting the economic share to 1.00 returns high-value reach to 44.1%, near cost-aware targeting. The operative mechanism is the reserved high-CLTV review slice. At this nominal point the graph residual is not an incremental routing signal. The residual does not identify model underestimation: positive-residual terciles have low mean probability (0.11-0.17) and observed churn slightly below pᵢ. Graph residual and model disagreement are distinct (correlation -0.29). Fig. 1 shows offers at high calibrated risk and human review almost entirely above the 75th CLTV percentile. Fig. 2 is the two-objective plane. High-value reach is 56.4% for economic shares 0.70-0.90 and falls to 44.1% at share 1.00. The 0.85 point is a convenient minority-review operating point, not a unique optimum. Across 11 scenario perturbations the high-value coverage advantage remained positive, but it was small when review cost was $40 or reviewer accuracy was 0.75. Point-estimate economic superiority changed sign in several scenarios.",
    )

    add_picture(doc, FIG1, 6.3)
    add_caption(
        doc,
        "Fig. 1. Knowledge-guided decision map. Horizontal axis: calibrated ensemble churn probability. Vertical axis: CLTV percentile. Marker size encodes graph motif risk. The dotted line is the high-value CLTV quantile (0.75).",
    )
    add_picture(doc, FIG2, 6.3)
    add_caption(
        doc,
        "Fig. 2. Budget-matched tradeoff between the outcome-anchored net-benefit proxy and high-value churn reach. The path traces alternative economic-allocation fractions.",
    )

    add_h1(doc, "5. Limitations")
    add_text(
        doc,
        "Business parameters are simulated. Reviewers are modelled by fixed sensitivity and specificity. One public dataset is used. High-value reach uses the same CLTV threshold that the policy gates on; the result shows that the intended coverage objective is achieved, not that a hidden value segment was discovered. The graph residual does not detect underestimation here and does not change the 0.85 allocation. The bootstrap conditions on frozen OOF scores. External validity requires another cohort or a later time split [1,2].",
        first_line=False,
    )

    add_h1(doc, "6. Conclusion")
    add_text(
        doc,
        "This short communication formulates churn intervention as a constrained decision problem: calibrated risk, a compact service-contract graph, CLTV as strategic priority, scarce review, and a shared budget [3,4,5,11,12,14]. On IBM Telco data, reserving a minority of that budget for high-CLTV review raises high-value churn reach by 12.6 percentage points versus cost-aware targeting (1000/1000 resamples). Overall churn reach stays comparable. The economic proxy is favourable in point estimate (968/1000 resamples) and remains interval-qualified. Ablation attributes the coverage gain to the CLTV review reserve, not to incremental graph-residual routing. That is a decision-support result, not a new classifier.",
        first_line=False,
    )
    add_h1(doc, "Data Availability")
    add_text(
        doc,
        "The IBM Telco Customer Churn workbook is a public telecommunications benchmark of the same class used in recent peer-reviewed churn studies [1,2,7]. Retention costs, success probabilities, reviewer parameters, margin, horizon, and budget are simulated scenario assumptions. Tables 1-2, Figs. 1-2, and the supplementary ablation, residual, disagreement, and sensitivity files are reproduced from frozen out-of-fold scores by letter_outputs.py and review_analyses.py in https://github.com/drsaikrishnathota1/kbs-churn-decision-intelligence, using results/experiment_a_frozen.",
        first_line=False,
    )

    add_h1(doc, "Declaration of Competing Interest")
    add_text(
        doc,
        "The author declares that there are no known competing financial interests or personal relationships that could have appeared to influence the work reported in this manuscript.",
        first_line=False,
    )

    add_h1(doc, "Declaration of Generative AI and AI-Assisted Technologies")
    add_text(
        doc,
        "During preparation of this manuscript, AI-assisted tools were used for language drafting, editing, organization, and manuscript formatting support. The author reviewed and edited the content and takes full responsibility for the final manuscript. The experimental data, tables, and figures were generated from executable simulation code and reproducible CSV outputs.",
        first_line=False,
    )

    add_h1(doc, "Funding")
    add_text(
        doc,
        "This research received no specific grant from any funding agency in the public, commercial, or not-for-profit sectors.",
        first_line=False,
    )

    add_h1(doc, "CRediT Authorship Contribution Statement")
    add_text(
        doc,
        "Sai Krishna Thota: Conceptualization, Methodology, Software, Validation, Formal analysis, Investigation, Data curation, Writing - original draft, Writing - review and editing, Visualization.",
        first_line=False,
    )

    add_h1(doc, "References")
    refs = [
        "[1] Manzoor, A., Qureshi, M. A., Kidney, E., & Longo, L. (2024). A review on machine learning methods for customer churn prediction and recommendations for business practitioners. IEEE Access, 12, 70434-70463. https://doi.org/10.1109/ACCESS.2024.3402092",
        "[2] Imani, M., Joudaki, M., Beikmohammadi, A., & Arabnia, H. R. (2025). Customer churn prediction: A systematic review of recent advances, trends, and challenges in machine learning and deep learning. Machine Learning and Knowledge Extraction, 7(3), 105. https://doi.org/10.3390/make7030105",
        "[3] Shahabikargar, M., Beheshti, A., Mansoor, W., Zhang, X., Foo, E. J., Jolfaei, A., Hanif, A., & Shabani, N. (2025). ChurnKB: A generative AI-enriched knowledge base for customer churn feature engineering. Algorithms, 18(4), 238. https://doi.org/10.3390/a18040238",
        "[4] Jiang, P., Liu, Z., Abedin, M. Z., Wang, J., Yang, W., & Dong, Q. (2024). Profit-driven weighted classifier with interpretable ability for customer churn prediction. Omega, 125, 103034. https://doi.org/10.1016/j.omega.2024.103034",
        "[5] Feng, Y., Yin, Y., Wang, D., Ignatius, J., Cheng, T. C. E., Marra, M., & Guo, Y. (2024). Enhancing e-commerce customer churn management with a profit- and AUC-focused prescriptive analytics approach. Journal of Business Research, 184, 114872. https://doi.org/10.1016/j.jbusres.2024.114872",
        "[6] Rao, C., Xu, Y., Xiao, X., Hu, F., & Goh, M. (2024). Imbalanced customer churn classification using a new multi-strategy collaborative processing method. Expert Systems with Applications, 247, 123251. https://doi.org/10.1016/j.eswa.2024.123251",
        "[7] Wang, C., Rao, C., Hu, F., Xiao, X., & Goh, M. (2024). Risk assessment of customer churn in telco using FCLCNN-LSTM model. Expert Systems with Applications, 248, 123352. https://doi.org/10.1016/j.eswa.2024.123352",
        "[8] Haddadi, S. J., Farshidvard, A., dos Santos Silva, F., dos Reis, J. C., & Reis, M. S. (2024). Customer churn prediction in imbalanced datasets with resampling methods: A comparative study. Expert Systems with Applications, 246, 123086. https://doi.org/10.1016/j.eswa.2023.123086",
        "[9] Joy, U. G., Hoque, K. E., Uddin, M. N., Chowdhury, L., & Park, S.-B. (2024). A big data-driven hybrid model for enhancing streaming service customer retention through churn prediction integrated with explainable AI. IEEE Access, 12, 69130-69150. https://doi.org/10.1109/ACCESS.2024.3401247",
        "[10] Poudel, S. S., Pokharel, S., & Timilsina, M. (2024). Explaining customer churn prediction in telecom industry using tabular machine learning models. Machine Learning with Applications, 17, 100567. https://doi.org/10.1016/j.mlwa.2024.100567",
        "[11] Văduva, A.-G., Oprea, S. V., Niculae, A.-M., Bâra, A., & Andreescu, A.-I. (2024). Improving churn detection in the banking sector: A machine learning approach with probability calibration techniques. Electronics, 13(22), 4527. https://doi.org/10.3390/electronics13224527",
        "[12] Kostopoulos, G., Davrazos, G., & Kotsiantis, S. (2024). Explainable artificial intelligence-based decision support systems: A recent review. Electronics, 13(14), 2842. https://doi.org/10.3390/electronics13142842",
        "[13] Guo, Z., Zhou, D., Yu, D., Zhou, Q., Wu, H., & Hao, A. (2024). An ontology-based method for knowledge reuse in the design for maintenance of complex products. Computers in Industry, 161, 104124. https://doi.org/10.1016/j.compind.2024.104124",
        "[14] Su, C., Jiang, Q., Han, Y., Wang, T., & He, Q. (2025). Knowledge graph-driven decision support for manufacturing process: A graph neural network-based knowledge reasoning approach. Advanced Engineering Informatics, 64, 103098. https://doi.org/10.1016/j.aei.2024.103098",
        "[15] Greif, L., Hauck, S., Kimmig, A., & Ovtcharova, J. (2025). A knowledge graph framework to support life cycle assessment for sustainable decision-making. Applied Sciences, 15(1), 175. https://doi.org/10.3390/app15010175",
        "[16] Park, C., Lee, H., Lee, S., & Jeong, O. (2025). Synergistic joint model of knowledge graph and LLM for enhancing XAI-based clinical decision support systems. Mathematics, 13(6), 949. https://doi.org/10.3390/math13060949",
        "[17] Marandi, S., Hu, Y.-S., & Modarres, M. (2025). Complex system diagnostics using a knowledge graph-informed and large language model-enhanced framework. Applied Sciences, 15(17), 9428. https://doi.org/10.3390/app15179428",
        "[18] Li, J., Qian, L., Liu, P., & Liu, T. (2024). Construction of legal knowledge graph based on knowledge-enhanced large language models. Information, 15(11), 666. https://doi.org/10.3390/info15110666",
        "[19] Yang, Y., Liu, X., Tu, X., Lu, Y., & Wang, Y. (2025). Automating the construction of environmental policy knowledge graph with large language models. Sustainability, 17(22), 10282. https://doi.org/10.3390/su172210282",
        "[20] Zhou, Q., Zhou, D., Wang, Y., Guo, Z., & Dai, C. (2024). Knowledge reuse for ontology modelling and application of maintenance motion state sequence. Journal of Industrial Information Integration, 41, 100659. https://doi.org/10.1016/j.jii.2024.100659",
    ]
    for r in refs:
        add_text(doc, r, first_line=False, after=6)

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
